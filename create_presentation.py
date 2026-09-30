from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "SF_Rent_Board_Final_Presentation.pptx"
ER_IMAGE = ROOT / "RentBoard_ER_Diagram.png"
NAVY = RGBColor(22, 43, 70)
TEAL = RGBColor(18, 130, 145)
ORANGE = RGBColor(232, 132, 58)
LIGHT = RGBColor(242, 247, 249)
DARK = RGBColor(38, 48, 58)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(95, 108, 120)


def text_box(slide, text, x, y, w, h, size=20, color=DARK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Aptos"
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    if align is not None:
        p.alignment = align
    return box


def header(slide, title, number):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.72))
    bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    text_box(slide, title, 0.55, 0.12, 10.5, 0.4, 26, WHITE, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.15), Inches(12.25), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(220, 228, 232); line.line.fill.background()
    text_box(slide, f"SF Rent Board Housing Inventory  |  DATA 201  |  {number}", 0.6, 7.2, 12, 0.2, 9, GRAY)


def slide_with_header(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = WHITE
    header(slide, title, len(prs.slides))
    return slide


def bullets(slide, items, x=0.9, y=1.25, w=11.5, h=5.5, size=20):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = DARK
        p.space_after = Pt(12)
    return box


def callout(slide, text, x, y, w, h, fill=LIGHT, line=TEAL, size=19):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill; shape.line.color.rgb = line
    tf = shape.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.text = text; p.alignment = PP_ALIGN.CENTER
    p.font.name = "Aptos"; p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = NAVY


prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

# Title slide with editable placeholders.
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.24), Inches(7.5))
stripe.fill.solid(); stripe.fill.fore_color.rgb = ORANGE; stripe.line.fill.background()
text_box(slide, "SF Rent Board Housing Inventory", 0.9, 1.05, 11.5, 0.8, 38, WHITE, True)
text_box(slide, "Cleaning assessment and 3NF relational database model", 0.95, 2.05, 11.2, 0.55, 23, RGBColor(205, 224, 232))
callout(slide, "Professor: [Professor Name]", 0.95, 3.65, 5.25, 0.75, RGBColor(36, 67, 94), TEAL, 18)
callout(slide, "Team: [Team Name]", 6.45, 3.65, 5.25, 0.75, RGBColor(36, 67, 94), ORANGE, 18)
text_box(slide, "Team members: [Add names here]", 0.98, 5.05, 11.2, 0.5, 20, WHITE)
text_box(slide, "DATA 201  |  Database design and normalization project", 0.98, 6.35, 11.3, 0.3, 13, RGBColor(180, 204, 215))

# Content slides.
slide = slide_with_header(prs, "Project overview")
bullets(slide, ["Build a reliable relational database from a large, messy housing-inventory export.", "Clean inconsistent labels, ranges, dates, utilities, and duplicate filings.", "Organize the data into a 3NF schema with clear primary keys and foreign keys.", "Validate that the database preserves the original 550,201 unit records."], size=21)
for i, label in enumerate(["550,201 records", "28 columns", "21 tables", "3NF design"]):
    callout(slide, label, 0.85 + i * 3.1, 5.95, 2.65, 0.65, RGBColor(231, 246, 247) if i % 2 == 0 else RGBColor(255, 244, 232), TEAL if i % 2 == 0 else ORANGE, 18)

slide = slide_with_header(prs, "Dataset selection and motivation")
bullets(slide, ["The data describes one rental-unit record from one annual filing.", "It is complex enough for database design because it includes categories, ranges, dates, utilities, history, and quality problems.", "The project supports questions about filing years, occupancy, rent bands, utilities, and data quality.", "Rent and square footage are bands, so midpoint assumptions must be stated clearly."], size=20)
callout(slide, "Goal: turn a wide flat file into a clean, auditable relational model", 1.2, 5.95, 10.9, 0.7, RGBColor(231, 246, 247), TEAL, 20)

slide = slide_with_header(prs, "What cleaning was needed")
bullets(slide, ["Bedroom and bathroom fields contained spelling variations, numeric text, categories, and unparseable labels.", "Monthly rent and square footage were stored as ranges rather than exact numeric values.", "Utility information was spread across checkboxes and a multi-valued free-text field.", "Invalid dates were preserved through quality flags instead of being silently dropped.", "Identical non-key records were tagged as duplicate groups."], size=18)

slide = slide_with_header(prs, "Entity-relationship diagram")
if ER_IMAGE.exists():
    slide.shapes.add_picture(str(ER_IMAGE), Inches(0.75), Inches(0.92), width=Inches(11.85), height=Inches(5.95))
else:
    callout(slide, "Insert RentBoard_ER_Diagram.png here", 2, 2.7, 9.2, 1, LIGHT, ORANGE, 24)

slide = slide_with_header(prs, "Normalization: 1NF, 2NF, and 3NF")
callout(slide, "1NF", 0.85, 1.3, 1.35, 0.65, RGBColor(231, 246, 247), TEAL, 22)
text_box(slide, "Utilities and occupancy history were moved into separate rows so values are atomic.", 2.5, 1.3, 9.9, 0.7, 18)
callout(slide, "2NF", 0.85, 2.7, 1.35, 0.65, RGBColor(255, 244, 232), ORANGE, 22)
text_box(slide, "unit_record uses the single-column key unique_id, so partial dependencies cannot occur.", 2.5, 2.7, 9.9, 0.7, 18)
callout(slide, "3NF", 0.85, 4.1, 1.35, 0.65, RGBColor(231, 246, 247), TEAL, 22)
text_box(slide, "Repeated descriptions and derived attributes were moved into lookup tables such as rent_band and occupancy_type.", 2.5, 4.1, 9.9, 0.85, 18)
callout(slide, "Result: less repetition, clearer relationships, easier validation", 1.25, 5.8, 10.8, 0.7, LIGHT, NAVY, 20)

slide = slide_with_header(prs, "Database design")
bullets(slide, ["unit_record is the central fact table and uses unique_id as its primary key.", "Lookup tables store reusable descriptions for filing cycles, occupancy, bedrooms, bathrooms, rent, and square footage.", "record_utility_included and record_quality_flag are junction tables because one record can have many related values.", "occupancy_history is a child table because one unit record may have multiple history entries.", "Primary keys and foreign keys make the relationships explicit."], size=19)

slide = slide_with_header(prs, "Loading and validation evidence")
for i, (number, label) in enumerate([("550,201", "unit records"), ("81,004", "history records"), ("700,552", "utility links"), ("0", "duplicate unique IDs")]):
    x = 0.8 + i * 3.1
    callout(slide, number, x, 1.35, 2.7, 0.9, RGBColor(231, 246, 247) if i % 2 == 0 else RGBColor(255, 244, 232), TEAL if i % 2 == 0 else ORANGE, 27)
    text_box(slide, label, x, 2.35, 2.7, 0.35, 15, GRAY, False, PP_ALIGN.CENTER)
bullets(slide, ["The loader stages TSV files, loads tables in foreign-key order, and records quality issues instead of dropping them.", "Validation checks row counts, uniqueness, foreign-key relationships, and occupancy-date representations."], x=1, y=3.55, w=11.2, h=1.6, size=20)

slide = slide_with_header(prs, "Analysis questions and findings")
bullets(slide, ["How many records were filed in each year?", "How are records distributed across occupancy types?", "Which rent and square-footage bands are most common?", "Which utilities are included in base rent?", "Which records contain invalid dates or other quality problems?"], size=21)
callout(slide, "Insert SQL result screenshots here", 2, 5.75, 9.3, 0.75, RGBColor(255, 244, 232), ORANGE, 21)

slide = slide_with_header(prs, "Limitations and conclusion")
bullets(slide, ["The supplied local flat CSV has blank block, neighborhood, and supervisor-district labels.", "Rent and square footage are ranges, not exact measurements.", "The source does not provide a stable unit identifier across years.", "The normalized database preserves the source records while making cleaning decisions visible and auditable."], size=19)
callout(slide, "Normalization makes a complex housing dataset easier to query, explain, and validate.", 1, 5.75, 11.2, 0.75, RGBColor(231, 246, 247), TEAL, 20)

prs.save(OUT)
print(OUT)
